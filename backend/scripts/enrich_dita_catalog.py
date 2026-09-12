"""Phase 1 of the DITA 1.3 BRDP catalog enrichment (docs request):
generates a REVIEW file only -- writes nothing to Postgres. Phase 2
(the real import into brdp_catalog with standard="Schematron 1.0 —
DITA") is a separate step that must not run until a human has reviewed
this file and explicitly approved it.

    cd backend && python scripts/enrich_dita_catalog.py

Reads catalog_sources/dita_1.3.xlsx, sheet "Dita BRDPs" -- a completely
different shape from import_brdp_catalog.py's "Auto-gen Decisions"
sheet, so that script is not reused here. Columns (confirmed against
the real file):

    A  BRDP reference/location
    B  BRDP unique identifier      (S1000D identifier, e.g. BRDP-S1-00001)
    C  BRDP title                  (S1000D title -- default Title candidate)
    D  S1000D BRDP definition
    E  ID                          (DITA identifier, BRDP-D1-NNNNN -- filter column)
    F  "Title" per the sheet's own header -- actually the DITA DEFINITION
    G  "Definition" per the sheet's own header -- actually the DITA PROPOSAL

F/G's header labels are swapped relative to their real content (confirmed
with the user) -- this script treats them by their real meaning, not
their literal header text. Only the 100 rows where column E is non-blank
are DITA-adapted BRDPs; the other ~327 rows are S1000D-only and skipped.

For each of the 100 rows, a real Mistral chat completion (same
endpoint/key resolution as app/services/embeddings.py, but a chat call,
not an embedding one -- there is no existing "plain chat completion from
the backend" helper to reuse, so this makes its own request rather than
force-fitting embeddings.py's endpoint-specific shape) decides:

  1. Title -- column C is the default candidate; the LLM judges per row
     whether it still fits the DITA content (F/G) and only writes a new
     one from F/G when it genuinely does not. Not a blanket rule.
  2. Definition -- synthesized from the real F content plus whatever
     from G helps understand WHAT decision is being made, dropping
     purely project-specific implementation detail. The proposal itself
     (G) is never carried into brdp_catalog's own definition column
     wholesale, and is not persisted at all in Phase 2.

The prompt is restricted to that single row's own text (columns A-G) --
no outside knowledge, nothing invented beyond what's given or reasonably
implicit in it.

Output: dita_enrichment_review.csv (in this scripts/ directory), one row
per DITA BRDP, for a human to read before Phase 2 exists at all.
"""
import asyncio
import csv
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import httpx
import openpyxl

from app.core.config import get_settings

SHEET_NAME = "Dita BRDPs"
XLSX_PATH = Path(__file__).resolve().parent.parent / "catalog_sources" / "dita_1.3.xlsx"
OUTPUT_PATH = Path(__file__).resolve().parent / "dita_enrichment_review.csv"

_DITA_ID_PATTERN = re.compile(r"^BRDP-D1-\d{5}$")

_SYSTEM_PROMPT = (
    "You curate a catalog of BRDP (Business Rule Decision Point) entries for DITA 1.3 "
    "authoring rules. You work strictly from the text given for ONE spreadsheet row. "
    "Never invent content that is not present, or reasonably implicit, in that text. "
    "Respond with strict JSON only, no markdown fences, no commentary: "
    '{"title": "...", "title_changed": true|false, "definition": "..."}'
)


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).replace("_x000D_", "\n").strip()


def _read_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet {SHEET_NAME!r} not found. Sheets in file: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)
    ws = wb[SHEET_NAME]

    rows = []
    seen_ids: dict[str, int] = {}  # identifier -> index into `rows` of its first occurrence
    for row_num in range(2, ws.max_row + 1):
        dita_id = ws.cell(row=row_num, column=5).value
        if dita_id is None or str(dita_id).strip() == "":
            continue
        dita_id = str(dita_id).strip()

        flags = []
        if not _DITA_ID_PATTERN.match(dita_id):
            flags.append(f"id_format_mismatch: expected BRDP-D1-NNNNN, got {dita_id!r}")
        if dita_id in seen_ids:
            first_row_num = rows[seen_ids[dita_id]]["row_num"]
            flags.append(f"duplicate_id: also used on row {first_row_num}")
            rows[seen_ids[dita_id]]["flags"].append(f"duplicate_id: also used on row {row_num}")
        else:
            seen_ids[dita_id] = len(rows)

        rows.append(
            {
                "row_num": row_num,
                "s1000d_reference": _clean(ws.cell(row=row_num, column=1).value),
                "s1000d_identifier": _clean(ws.cell(row=row_num, column=2).value),
                "s1000d_title": _clean(ws.cell(row=row_num, column=3).value),
                "s1000d_definition": _clean(ws.cell(row=row_num, column=4).value),
                "identifier": dita_id,
                "original_definition": _clean(ws.cell(row=row_num, column=6).value),  # col F, real meaning
                "original_proposal": _clean(ws.cell(row=row_num, column=7).value),  # col G, real meaning
                "flags": flags,
            }
        )
    return rows


def _build_user_prompt(row: dict) -> str:
    return (
        "Spreadsheet row (columns A-G):\n"
        f"A (BRDP reference/location): {row['s1000d_reference']}\n"
        f"B (S1000D BRDP identifier): {row['s1000d_identifier']}\n"
        f"C (S1000D BRDP title): {row['s1000d_title']}\n"
        f"D (S1000D BRDP definition): {row['s1000d_definition']}\n"
        f"E (DITA BRDP identifier): {row['identifier']}\n"
        f"F (DITA definition -- the sheet's own header calls this column \"Title\", "
        f"but its real content is the definition): {row['original_definition']}\n"
        f"G (DITA proposal -- the sheet's own header calls this column \"Definition\", "
        f"but its real content is the proposal): {row['original_proposal']}\n"
        "\n"
        "Task 1 -- Title: column C is the default candidate title for this DITA BRDP "
        "catalog entry, since it is what S1000D uses for the same underlying rule. "
        "Judge whether it still fits the DITA content in F/G. If it fits, reuse it "
        "as-is (title_changed=false). If it does not genuinely fit (e.g. it names an "
        "S1000D-specific mechanism that F/G is not actually about), write a new, short "
        "title drawn only from F/G (title_changed=true).\n"
        "Task 2 -- Definition: synthesize the final definition by combining the real "
        "content of F with whatever from G is relevant to understanding WHAT decision "
        "is being made -- leave out purely project-specific implementation detail from "
        "G that does not help understand the decision itself."
    )


def _extract_json(text: str) -> dict:
    """Mirrors the frontend's own tolerance for LLM output that isn't pure
    JSON (extractBRDPs.js does the same) -- grabs the first {...} block
    rather than assuming response_format is honored by whatever endpoint
    ACTIVE_LLM_PROVIDER points at (may be a private, non-OpenAI-compatible
    deployment, docs/v2 §3 point 6).
    """
    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1 or end < start:
        raise ValueError(f"No JSON object found in LLM response: {text!r}")
    return json.loads(text[start : end + 1])


async def _call_llm(client: httpx.AsyncClient, endpoint: str, api_key: str, model: str, row: dict) -> dict:
    payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user", "content": _build_user_prompt(row)},
        ],
    }
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"}
    response = await client.post(endpoint, headers=headers, json=payload)
    response.raise_for_status()
    content = response.json()["choices"][0]["message"]["content"]
    parsed = _extract_json(content)
    return {
        "title": str(parsed["title"]).strip(),
        "title_changed": bool(parsed.get("title_changed", False)),
        "definition": str(parsed["definition"]).strip(),
    }


async def main() -> None:
    settings = get_settings()
    if settings.active_llm_provider != "mistral":
        print(
            f"ACTIVE_LLM_PROVIDER is {settings.active_llm_provider!r}, not 'mistral' -- "
            "this script only targets Mistral (same constraint as embeddings.py).",
            file=sys.stderr,
        )
        sys.exit(1)
    if not settings.mistral_endpoint or not settings.mistral_api_key:
        print("Mistral chat endpoint/API key is not fully configured on the server.", file=sys.stderr)
        sys.exit(1)
    if not XLSX_PATH.is_file():
        print(f"File not found: {XLSX_PATH}", file=sys.stderr)
        sys.exit(1)

    rows = _read_rows()
    print(f"Found {len(rows)} DITA-adapted BRDPs (column E non-blank) out of the full sheet.")

    results = []
    async with httpx.AsyncClient(timeout=60.0) as client:
        for i, row in enumerate(rows, start=1):
            print(f"[{i}/{len(rows)}] {row['identifier']} ...", end=" ", flush=True)
            try:
                llm_out = await _call_llm(
                    client, settings.mistral_endpoint, settings.mistral_api_key, settings.mistral_model, row
                )
                print("ok" + (" (title changed)" if llm_out["title_changed"] else ""))
            except Exception as err:  # noqa: BLE001 -- one row's failure must not lose the batch
                print(f"FAILED: {err}")
                row["flags"].append(f"llm_call_failed: {err}")
                llm_out = {"title": "", "title_changed": False, "definition": ""}
            results.append({**row, **llm_out})

    with OUTPUT_PATH.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "identifier",
                "s1000d_identifier",
                "s1000d_title",
                "original_definition",
                "original_proposal",
                "llm_title",
                "llm_title_changed",
                "llm_definition",
                "flags",
            ],
        )
        writer.writeheader()
        for r in results:
            writer.writerow(
                {
                    "identifier": r["identifier"],
                    "s1000d_identifier": r["s1000d_identifier"],
                    "s1000d_title": r["s1000d_title"],
                    "original_definition": r["original_definition"],
                    "original_proposal": r["original_proposal"],
                    "llm_title": r["title"],
                    "llm_title_changed": r["title_changed"],
                    "llm_definition": r["definition"],
                    "flags": "; ".join(r["flags"]),
                }
            )

    failed = sum(1 for r in results if any("llm_call_failed" in f for f in r["flags"]))
    flagged = sum(1 for r in results if r["flags"])
    print(
        f"\nWrote {len(results)} rows to {OUTPUT_PATH} "
        f"({failed} LLM call failures, {flagged} rows with at least one flag). "
        "Nothing was written to Postgres -- review this file before Phase 2."
    )


if __name__ == "__main__":
    asyncio.run(main())
