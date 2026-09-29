"""Excel in and out, on the server (openpyxl) -- replaces SheetJS (xlsx) in
the browser.

Why: xlsx@0.18.5 has two high-severity advisories when READING a malicious
file (prototype pollution, ReDoS), and the official 0.20.3 altered rules
with XML entities on the way out (Juanma's Windows test: rules with &lt;,
&gt; or &#NNN; came back changed, and the ones with &lt; were then rejected
on reimport). openpyxl was already the backend's Excel library
(rule_templates.py, scripts/import_brdp_catalog.py, enrich_dita_catalog.py).

Three operations, all keeping the shape the frontend used before:
  parse_import_file(data, filename)  -> the rows/errors importFromExcel gave
  build_export_workbook(rows)        -> the file exportToExcel wrote
  build_generic_template()           -> the file generateTemplate wrote

Cell text is exact both ways: XML entities inside a rule (&lt;, &#237;...)
are plain characters of the cell text and are never decoded or encoded;
Excel's own escapes (_xHHHH_, e.g. _x000d_ for a CR) are undone on read
exactly like rule_templates.py (cell_text(), which that module now
imports), and written back on export only where Excel needs them.

Reading is defensive: only .xlsx (a zip with a workbook in it), a maximum
file size, a maximum uncompressed size (zip bombs), a maximum number of
rows, openpyxl in read_only mode with data_only=True (a formula is never
evaluated: its cached value is read, or nothing when there is none), and
defusedxml installed so openpyxl parses the sheet XML without entity
expansion. A file that fails any of these raises ExcelFileError, which the
route turns into a 422 with the reason (HR7) -- nothing is imported.
"""

import io
import json
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.utils.escape import unescape

from app.core.config import get_settings

# Column header -> import row key. Same 7 columns, same order, as the export
# (ID/Title/Definition/Proposal/Proposal Status/Rule Status/Rule).
IMPORT_FIELD_MAP = {
    "ID": "identifier",
    "Title": "title",
    "Definition": "definition",
    "Proposal": "proposal",
    "Proposal Status": "proposal_status",
    "Rule Status": "rule_status",
    "Rule": "rule",
}

# Export column header -> key of the rows the frontend builds with
# brdpToExportRow() (ProjectConfigPage.jsx), and the column width (Excel
# character units, the same "wch" values SheetJS was given).
EXPORT_COLUMNS = [
    ("ID", "id", 20),
    ("Title", "title", 30),
    ("Definition", "definition", 40),
    ("Proposal", "proposal", 40),
    ("Proposal Status", "proposalStatus", 16),
    ("Rule Status", "ruleStatus", 14),
    ("Rule", "rule", 60),
]
SHEET_NAME = "BRDPs"

# Excel's hard limit on the text of one cell. Over it Excel refuses (or
# truncates) the cell -- the export is refused instead, naming the BRDP and
# the field, never cut (HR6/HR7).
EXCEL_CELL_CHAR_LIMIT = 32767

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_GENERIC_TEMPLATE_FILE = Path(__file__).resolve().parent / "generic_template_rows.json"

# Excel's escape for a character in cell text: _xHHHH_. On write, a literal
# "_xHHHH_" already in the text must itself be escaped (its "_" as _x005F_),
# or Excel -- and our own reader -- would decode it into a character.
_EXCEL_ESCAPE_RE = re.compile(r"_x[0-9A-Fa-f]{4}_")
# Control characters XML 1.0 cannot carry at all (tab, LF and CR can).
_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ExcelFileError(ValueError):
    """The uploaded file cannot be read safely: not an .xlsx, corrupt, or
    over one of the limits. The message is the reason shown to the user."""


class ExportCellTooLarge(ValueError):
    """At least one export cell is over EXCEL_CELL_CHAR_LIMIT. `cells` names
    each one: {"id", "field", "length"}."""

    def __init__(self, cells: list[dict]):
        super().__init__(f"{len(cells)} cell(s) over Excel's {EXCEL_CELL_CHAR_LIMIT}-character limit")
        self.cells = cells


# ─── Reading ──────────────────────────────────────────────────────────────


def cell_text(value) -> str:
    """The text of a cell as a user reads it in Excel. Strings: Excel's own
    escapes decoded (_x000d_ -> CR) and line breaks normalized to "\\n" (a
    CRLF from a pasted text, or the CR of a decoded _x000d_, never reaches
    the stored text). Nothing else is touched -- "&lt;" stays "&lt;".
    Numbers as Excel shows them (5, not 5.0); booleans TRUE/FALSE; dates in
    ISO form; empty -> "".
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if not isinstance(value, str):
        return str(value)
    return unescape(value).replace("\r\n", "\n").replace("\r", "\n")


def _check_zip(data: bytes) -> None:
    settings = get_settings()
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExcelFileError("The file is not an .xlsx workbook (only .xlsx files can be imported).")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {e.filename for e in entries}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ExcelFileError("The file is not an .xlsx workbook (only .xlsx files can be imported).")
            total = sum(e.file_size for e in entries)
    except zipfile.BadZipFile as exc:
        raise ExcelFileError(f"The file is corrupt and cannot be read: {exc}") from exc
    if total > settings.excel_import_max_uncompressed_bytes:
        raise ExcelFileError(
            f"The workbook expands to {total} bytes, over the {settings.excel_import_max_uncompressed_bytes}-byte limit."
        )


def parse_import_file(data: bytes, filename: str | None) -> dict:
    """Rows of the first sheet, keyed like ImportRowIn, with their real Excel
    row number (the header is row 1). Returns {"rows", "errors"}: errors are
    the same messages importFromExcel gave for a readable workbook that has
    nothing to import (no data rows, missing columns) -- the page shows them
    and imports nothing. A file that cannot be read safely raises
    ExcelFileError instead.

    Header cells are compared trimmed; extra columns are ignored. Rows with
    every mapped cell empty are skipped (their row number is simply not
    used). Every value arrives as text (see cell_text).
    """
    settings = get_settings()
    name = (filename or "").strip()
    if not name.lower().endswith(".xlsx"):
        raise ExcelFileError(f"Only .xlsx files can be imported (got '{name or 'no file name'}').")
    if not data:
        raise ExcelFileError("The file is empty.")
    if len(data) > settings.excel_import_max_bytes:
        raise ExcelFileError(f"The file is {len(data)} bytes, over the {settings.excel_import_max_bytes}-byte limit.")
    _check_zip(data)

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for a broken package
        raise ExcelFileError(f"The file is corrupt and cannot be read: {exc}") from exc
    try:
        if not workbook.worksheets:
            return {"rows": [], "errors": ["Excel file is empty"]}
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter, None)
        except Exception as exc:
            raise ExcelFileError(f"The file is corrupt and cannot be read: {exc}") from exc
        if header_row is None:
            return {"rows": [], "errors": ["No data rows found in Excel file"]}
        header = [cell_text(h).strip() for h in header_row]
        missing = [col for col in IMPORT_FIELD_MAP if col not in header]
        if missing:
            return {"rows": [], "errors": [f"Missing required columns: {', '.join(missing)}"]}
        index = {col: header.index(col) for col in IMPORT_FIELD_MAP}

        rows: list[dict] = []
        try:
            for row_number, raw in enumerate(rows_iter, start=2):
                values = {
                    key: cell_text(raw[index[col]]) if index[col] < len(raw) else ""
                    for col, key in IMPORT_FIELD_MAP.items()
                }
                if not any(v.strip() for v in values.values()):
                    continue
                if len(rows) >= settings.excel_import_max_rows:
                    raise ExcelFileError(
                        f"The file has more than {settings.excel_import_max_rows} data rows; split it into smaller files."
                    )
                rows.append({"row_number": row_number, **values})
        except ExcelFileError:
            raise
        except Exception as exc:
            raise ExcelFileError(f"The file is corrupt and cannot be read: {exc}") from exc
    finally:
        workbook.close()

    if not rows:
        return {"rows": [], "errors": ["No data rows found in Excel file"]}
    return {"rows": rows, "errors": []}


# ─── Writing ──────────────────────────────────────────────────────────────


def _excel_text(value: str) -> str:
    """Cell text as it must be written so Excel (and cell_text) read back
    exactly `value`: a literal _xHHHH_ has its "_" escaped, and a control
    character XML cannot carry is written as Excel's _xHHHH_ escape. Tab,
    LF and CR are written as they are -- no _x000D_ is ever added."""
    text = _EXCEL_ESCAPE_RE.sub(lambda m: "_x005F" + m.group(0), value)
    return _XML_ILLEGAL_RE.sub(lambda m: f"_x{ord(m.group(0)):04X}_", text)


def _workbook_bytes(rows: list[dict]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append([header for header, _key, _width in EXPORT_COLUMNS])
    for row in rows:
        sheet.append([_excel_text(str(row.get(key) or "")) for _header, key, _width in EXPORT_COLUMNS])
    for position, (_header, _key, width) in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(position)].width = width
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def oversized_export_cells(rows: list[dict]) -> list[dict]:
    """Every cell over Excel's limit: {"id", "field" (the column header),
    "length"}, in row order."""
    cells = []
    for row in rows:
        for header, key, _width in EXPORT_COLUMNS:
            length = len(str(row.get(key) or ""))
            if length > EXCEL_CELL_CHAR_LIMIT:
                cells.append({"id": str(row.get("id") or ""), "field": header, "length": length})
    return cells


def build_export_workbook(rows: list[dict]) -> bytes:
    """The export file: sheet "BRDPs", the 7 columns in order, one row per
    BRDP exactly as the frontend built it (the server never recomputes rule
    states). Raises ExportCellTooLarge -- nothing is cut."""
    oversized = oversized_export_cells(rows)
    if oversized:
        raise ExportCellTooLarge(oversized)
    return _workbook_bytes(rows)


def generic_template_rows() -> list[dict]:
    return json.loads(_GENERIC_TEMPLATE_FILE.read_text(encoding="utf-8"))["rows"]


def build_generic_template() -> bytes:
    """The generic template for a standard with no curated one (S1000D
    5.0/6.0): the same 10 rows generateTemplate() built, Rule Status To Do
    and an empty Rule on each."""
    return _workbook_bytes(generic_template_rows())
