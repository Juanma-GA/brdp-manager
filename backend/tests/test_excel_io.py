"""Excel on the server (app/services/excel_io.py) -- replaces SheetJS.

The edge-case table of the task: a rule with &lt;parameter&gt; identical
after export and reimport; line breaks kept without _x000D_ or \\r added; a
file that is not .xlsx, is corrupt or too large -> 422 and nothing imported;
a cell over 32,767 characters in the export -> 422 naming the BRDP and the
field, nothing cut; a workbook with formulas -> the value is read, nothing
is evaluated. Plus the five curated templates of public/ read -> exported ->
read again cell for cell, and the generic template.
"""
import io
import json
import os
import re
import subprocess
import sys
import uuid
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.core.config import get_settings
from app.core.security import create_access_token, hash_password
from app.db.base import async_session_factory
from app.models import BRDP, Project, User, UserProjectRole
from app.services.excel_io import (
    EXCEL_CELL_CHAR_LIMIT,
    EXPORT_COLUMNS,
    ExcelFileError,
    ExportCellTooLarge,
    build_export_workbook,
    build_generic_template,
    cell_text,
    parse_import_file,
)
from app.services.rule_templates import CURATED_TEMPLATE_BY_STANDARD, PUBLIC_DIR

HEADERS = [header for header, _key, _width in EXPORT_COLUMNS]

# A rule with XML entities (&lt;x&gt;, a bare &gt;, a character reference),
# several lines and both kinds of quote -- the cell text Juanma's Windows
# test saw altered by SheetJS 0.20.3.
ENTITY_RULE = (
    '<structureObjectRule id="BRDP-XL-00001">\n'
    '  <objectPath allowedObjectFlag="0">//para[. = \'&lt;parameter&gt;\']</objectPath>\n'
    "  <objectUse>No &lt;x&gt; element, no a &gt; b, ni acción con tilde (&#237;) ni \"comillas\".</objectUse>\n"
    "</structureObjectRule>"
)


def _export_row(identifier="BRDP-XL-00001", rule=ENTITY_RULE, **extra):
    row = {
        "id": identifier,
        "title": "Título con \"comillas\" y 'simples'",
        "definition": "Línea 1\nLínea 2\n\nLínea 4",
        "proposal": "Proposal with <table> and & and _x000D_ written literally",
        "proposalStatus": "Validated",
        "ruleStatus": "Verified",
        "rule": rule,
    }
    row.update(extra)
    return row


def _import_row_values(export_row):
    return {
        "identifier": export_row["id"],
        "title": export_row["title"],
        "definition": export_row["definition"],
        "proposal": export_row["proposal"],
        "proposal_status": export_row["proposalStatus"],
        "rule_status": export_row["ruleStatus"],
        "rule": export_row["rule"],
    }


def _workbook(rows, header=HEADERS, sheet_title="BRDPs"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    if header is not None:
        sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _replace_in_sheet(data: bytes, replace) -> bytes:
    """Rewrites xl/worksheets/sheet1.xml -- to build what openpyxl never
    writes itself (Excel's _x000D_ escapes, cached formula values)."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                changed = replace(content.decode("utf-8"))
                assert changed != content.decode("utf-8"), "the sheet XML was not rewritten"
                content = changed.encode("utf-8")
            target.writestr(item, content)
    return out.getvalue()


# ─── cell_text ────────────────────────────────────────────────────────────


def test_cell_text_values():
    assert cell_text(None) == ""
    assert cell_text("a_x000d_\nb") == "a\nb"
    assert cell_text("a\r\nb\rc") == "a\nb\nc"
    assert cell_text("&lt;x&gt; &#237;") == "&lt;x&gt; &#237;"
    assert cell_text(5) == "5"
    assert cell_text(5.0) == "5"
    assert cell_text(2.5) == "2.5"
    assert cell_text(True) == "TRUE"


# ─── Reading ──────────────────────────────────────────────────────────────


def test_rule_with_entities_is_identical_after_export_and_reimport():
    rows = [_export_row(), _export_row("BRDP-XL-00002", rule="")]
    parsed = parse_import_file(build_export_workbook(rows), "brdps-export.xlsx")
    assert parsed["errors"] == []
    assert [r["row_number"] for r in parsed["rows"]] == [2, 3]
    for got, sent in zip(parsed["rows"], rows):
        assert {k: v for k, v in got.items() if k != "row_number"} == _import_row_values(sent)
    assert "&lt;parameter&gt;" in parsed["rows"][0]["rule"]


def test_line_breaks_are_kept_without_x000d_or_cr():
    parsed = parse_import_file(build_export_workbook([_export_row()]), "x.xlsx")
    definition = parsed["rows"][0]["definition"]
    assert definition == "Línea 1\nLínea 2\n\nLínea 4"
    assert "\r" not in definition and "_x000" not in definition.lower()
    # The literal text "_x000D_" typed in a cell survives as text.
    assert parsed["rows"][0]["proposal"].endswith("_x000D_ written literally")


def test_excel_own_escapes_and_crlf_are_normalized():
    # What Excel writes for a pasted CRLF: _x000D_ followed by a real LF.
    data = _workbook([["BRDP-XL-1", "T", "line1\nline2", "", "Pending", "To Do", ""]])
    data = _replace_in_sheet(data, lambda xml: xml.replace("line1\nline2", "line1_x000D_\nline2"))
    assert parse_import_file(data, "x.xlsx")["rows"][0]["definition"] == "line1\nline2"


def test_control_characters_and_literal_escapes_roundtrip():
    row = _export_row(proposal="bell\x07 and _x0041_ literally", rule="")
    parsed = parse_import_file(build_export_workbook([row]), "x.xlsx")
    assert parsed["rows"][0]["proposal"] == "bell\x07 and _x0041_ literally"


def test_formulas_are_never_evaluated():
    data = _workbook([["BRDP-XL-1", "=1+1", '=HYPERLINK("http://example.com","x")', "", "Pending", "To Do", ""]])
    # Title with a cached value (what Excel stores next to a formula): the
    # value is read. Definition without one (openpyxl never computes it):
    # nothing, and the formula text never reaches the row either.
    data = _replace_in_sheet(data, lambda xml: re.sub(r"(<c r=\"B2\"[^>]*>)<f>1\+1</f>", r"\1<f>1+1</f><v>2</v>", xml))
    row = parse_import_file(data, "x.xlsx")["rows"][0]
    assert row["title"] == "2"
    assert row["definition"] == ""
    assert "HYPERLINK" not in str(row)


def test_numbers_blank_rows_and_extra_columns():
    header = ["Notes", *HEADERS]
    data = _workbook(
        [
            ["x", "BRDP-XL-1", "T", "D", "", "Pending", "To Do", ""],
            [None] * 8,
            ["y", 42, 3.0, "D", "", "Validated", "To Do", ""],
        ],
        header=[f"  {h} " for h in header],
    )
    parsed = parse_import_file(data, "x.xlsx")
    assert [r["row_number"] for r in parsed["rows"]] == [2, 4]
    assert parsed["rows"][1]["identifier"] == "42" and parsed["rows"][1]["title"] == "3"


def test_missing_columns_and_no_rows_are_errors_not_422():
    missing = parse_import_file(_workbook([["a", "b"]], header=["ID", "Title"]), "x.xlsx")
    assert missing["rows"] == []
    assert missing["errors"] == ["Missing required columns: Definition, Proposal, Proposal Status, Rule Status, Rule"]
    assert parse_import_file(_workbook([]), "x.xlsx") == {"rows": [], "errors": ["No data rows found in Excel file"]}
    assert parse_import_file(_workbook([], header=None), "x.xlsx") == {"rows": [], "errors": ["No data rows found in Excel file"]}


@pytest.mark.parametrize(
    ("data", "filename", "reason"),
    [
        (b"ID,Title\n1,2\n", "brdps.csv", "Only .xlsx"),
        (b"\xd0\xcf\x11\xe0 old binary workbook", "brdps.xls", "Only .xlsx"),
        (b"just text", "brdps.xlsx", "not an .xlsx workbook"),
        (b"PK\x03\x04 truncated zip", "brdps.xlsx", "not an .xlsx workbook"),
        (b"", "brdps.xlsx", "empty"),
    ],
)
def test_not_an_xlsx_is_refused(data, filename, reason):
    with pytest.raises(ExcelFileError, match=reason):
        parse_import_file(data, filename)


def test_a_zip_that_is_not_a_workbook_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hello")
    with pytest.raises(ExcelFileError, match="not an .xlsx workbook"):
        parse_import_file(out.getvalue(), "x.xlsx")


def test_a_corrupt_workbook_is_refused():
    data = _replace_in_sheet(_workbook([["BRDP-XL-1", "T", "D", "", "Pending", "To Do", ""]]), lambda xml: xml[: len(xml) // 2])
    with pytest.raises(ExcelFileError, match="corrupt"):
        parse_import_file(data, "x.xlsx")


def test_size_rows_and_uncompressed_limits(monkeypatch):
    settings = get_settings()
    data = _workbook([[f"BRDP-XL-{i}", "T", "D", "", "Pending", "To Do", ""] for i in range(5)])
    monkeypatch.setattr(settings, "excel_import_max_bytes", len(data) - 1)
    with pytest.raises(ExcelFileError, match="byte limit"):
        parse_import_file(data, "x.xlsx")
    monkeypatch.setattr(settings, "excel_import_max_bytes", 10 * 1024 * 1024)
    monkeypatch.setattr(settings, "excel_import_max_rows", 4)
    with pytest.raises(ExcelFileError, match="more than 4 data rows"):
        parse_import_file(data, "x.xlsx")
    monkeypatch.setattr(settings, "excel_import_max_rows", 5)
    assert len(parse_import_file(data, "x.xlsx")["rows"]) == 5
    monkeypatch.setattr(settings, "excel_import_max_uncompressed_bytes", 1000)
    with pytest.raises(ExcelFileError, match="expands to"):
        parse_import_file(data, "x.xlsx")


# ─── Writing ──────────────────────────────────────────────────────────────


def test_text_that_looks_like_a_formula_stays_text():
    # openpyxl turns a string starting with "=" into a formula: read back
    # with data_only=True it came back empty (data lost), and Excel would
    # evaluate it when the export is opened. Every exported cell is text.
    row = _export_row(
        "BRDP-XL-FORMULA",
        title="=1+1",
        definition='=HYPERLINK("http://x","y")',
        proposal="+x",
        proposalStatus="-y",
        ruleStatus="@z",
        rule="=SUM(1,2)",
    )
    data = build_export_workbook([row])
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert [(c.value, c.data_type) for c in sheet[2]] == [
        ("BRDP-XL-FORMULA", "s"),
        ("=1+1", "s"),
        ('=HYPERLINK("http://x","y")', "s"),
        ("+x", "s"),
        ("-y", "s"),
        ("@z", "s"),
        ("=SUM(1,2)", "s"),
    ]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "<f>" not in archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    parsed = parse_import_file(data, "brdps-export.xlsx")["rows"][0]
    assert {k: v for k, v in parsed.items() if k != "row_number"} == _import_row_values(row)
    # The generic template goes through the same writer.
    template = openpyxl.load_workbook(io.BytesIO(build_generic_template())).active
    assert all(c.data_type in ("s", "inlineStr") for r in template.iter_rows() for c in r if c.value is not None)


def test_export_layout():
    workbook = openpyxl.load_workbook(io.BytesIO(build_export_workbook([_export_row()])))
    assert workbook.sheetnames == ["BRDPs"]
    sheet = workbook["BRDPs"]
    assert [c.value for c in sheet[1]] == HEADERS
    assert [sheet.column_dimensions[letter].width for letter in "ABCDEFG"] == [20, 30, 40, 40, 16, 14, 60]


def test_export_cell_over_the_limit_is_refused_and_never_cut():
    at_limit = _export_row("BRDP-XL-OK", rule="x" * EXCEL_CELL_CHAR_LIMIT)
    parsed = parse_import_file(build_export_workbook([at_limit]), "x.xlsx")
    assert len(parsed["rows"][0]["rule"]) == EXCEL_CELL_CHAR_LIMIT
    rows = [at_limit, _export_row("BRDP-XL-BIG", rule="x" * (EXCEL_CELL_CHAR_LIMIT + 1), title="t" * 40000)]
    with pytest.raises(ExportCellTooLarge) as caught:
        build_export_workbook(rows)
    assert caught.value.cells == [
        {"id": "BRDP-XL-BIG", "field": "Title", "length": 40000},
        {"id": "BRDP-XL-BIG", "field": "Rule", "length": EXCEL_CELL_CHAR_LIMIT + 1},
    ]


def test_generic_template():
    parsed = parse_import_file(build_generic_template(), "brdp-template.xlsx")
    assert parsed["errors"] == []
    assert len(parsed["rows"]) == 10
    assert parsed["rows"][0]["identifier"] == "BRDP-S1-00001"
    assert {r["rule_status"] for r in parsed["rows"]} == {"To Do"}
    assert {r["rule"] for r in parsed["rows"]} == {""}


@pytest.mark.parametrize("standard", sorted(CURATED_TEMPLATE_BY_STANDARD))
def test_curated_templates_roundtrip_cell_for_cell(standard):
    path = Path(PUBLIC_DIR) / CURATED_TEMPLATE_BY_STANDARD[standard]
    first = parse_import_file(path.read_bytes(), path.name)
    assert first["errors"] == [] and len(first["rows"]) == 10
    assert all(r["rule"].strip() for r in first["rows"])
    exported = build_export_workbook(
        [
            {
                "id": r["identifier"],
                "title": r["title"],
                "definition": r["definition"],
                "proposal": r["proposal"],
                "proposalStatus": r["proposal_status"],
                "ruleStatus": r["rule_status"],
                "rule": r["rule"],
            }
            for r in first["rows"]
        ]
    )
    again = parse_import_file(exported, "brdps-export.xlsx")
    assert [{k: v for k, v in r.items() if k != "row_number"} for r in again["rows"]] == [
        {k: v for k, v in r.items() if k != "row_number"} for r in first["rows"]
    ]


# ─── Endpoints ────────────────────────────────────────────────────────────


@pytest.fixture
async def project_and_users():
    async with async_session_factory() as session:
        project = Project(name=f"Excel Test Project {uuid.uuid4()}", standard="S1000D 4.2")
        editor = User(email=f"excel-editor-{uuid.uuid4()}@example.com", password_hash=hash_password("x-password-1"), display_name="E", global_role="user")
        viewer = User(email=f"excel-viewer-{uuid.uuid4()}@example.com", password_hash=hash_password("x-password-1"), display_name="V", global_role="user")
        session.add_all([project, editor, viewer])
        await session.flush()
        session.add_all(
            [
                UserProjectRole(user_id=editor.id, project_id=project.id, role="editor"),
                UserProjectRole(user_id=viewer.id, project_id=project.id, role="viewer"),
            ]
        )
        await session.commit()
        ids = (project.id, editor.id, viewer.id)
    project_id, editor_id, viewer_id = ids
    yield (
        project_id,
        {"Authorization": f"Bearer {create_access_token(editor_id)}"},
        {"Authorization": f"Bearer {create_access_token(viewer_id)}"},
    )
    async with async_session_factory() as session:
        for model, key in ((Project, project_id), (User, editor_id), (User, viewer_id)):
            obj = await session.get(model, key)
            if obj is not None:
                await session.delete(obj)
        await session.commit()


def _upload(data, filename="brdps.xlsx"):
    return {"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


async def test_parse_endpoint_returns_rows_that_analyze_accepts(client, project_and_users):
    project_id, editor, _viewer = project_and_users
    resp = await client.post(
        f"/api/projects/{project_id}/brdps/import/parse", files=_upload(build_export_workbook([_export_row()])), headers=editor
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["errors"] == [] and body["rows"][0]["rule"] == ENTITY_RULE
    analyzed = await client.post(f"/api/projects/{project_id}/brdps/import/analyze", json={"rows": body["rows"]}, headers=editor)
    assert analyzed.status_code == 200
    assert analyzed.json()["results"][0]["outcome"] == "ok"


async def test_parse_endpoint_readable_but_empty_gives_errors(client, project_and_users):
    project_id, editor, _viewer = project_and_users
    resp = await client.post(f"/api/projects/{project_id}/brdps/import/parse", files=_upload(_workbook([])), headers=editor)
    assert resp.status_code == 200
    assert resp.json() == {"rows": [], "errors": ["No data rows found in Excel file"]}


@pytest.mark.parametrize(
    ("data", "filename", "reason"),
    [
        (b"not a workbook", "brdps.xls", "Only .xlsx"),
        (b"not a workbook", "brdps.xlsx", "not an .xlsx workbook"),
    ],
)
async def test_parse_endpoint_refuses_bad_files_with_422(client, project_and_users, data, filename, reason):
    project_id, editor, _viewer = project_and_users
    resp = await client.post(f"/api/projects/{project_id}/brdps/import/parse", files=_upload(data, filename), headers=editor)
    assert resp.status_code == 422
    assert reason in resp.json()["detail"]
    async with async_session_factory() as session:
        from sqlalchemy import func, select

        assert (await session.execute(select(func.count()).select_from(BRDP).where(BRDP.project_id == project_id))).scalar_one() == 0


async def test_parse_endpoint_refuses_too_large_upload(client, project_and_users, monkeypatch):
    project_id, editor, _viewer = project_and_users
    data = _workbook([["BRDP-XL-1", "T", "D", "", "Pending", "To Do", ""]])
    monkeypatch.setattr(get_settings(), "excel_import_max_bytes", 100)
    resp = await client.post(f"/api/projects/{project_id}/brdps/import/parse", files=_upload(data), headers=editor)
    assert resp.status_code == 422
    assert "over the 100-byte limit" in resp.json()["detail"]


async def test_parse_is_editor_only_export_is_viewer(client, project_and_users):
    project_id, _editor, viewer = project_and_users
    parse = await client.post(f"/api/projects/{project_id}/brdps/import/parse", files=_upload(_workbook([])), headers=viewer)
    assert parse.status_code == 403
    export = await client.post(f"/api/projects/{project_id}/export.xlsx", json={"rows": [_export_row()]}, headers=viewer)
    assert export.status_code == 200
    assert export.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert export.headers["content-disposition"] == 'attachment; filename="brdps-export.xlsx"'
    assert parse_import_file(export.content, "brdps-export.xlsx")["rows"][0]["rule"] == ENTITY_RULE
    outsider = await client.post(f"/api/projects/{uuid.uuid4()}/export.xlsx", json={"rows": []}, headers=viewer)
    assert outsider.status_code == 403


async def test_export_endpoint_names_the_brdp_and_field_over_the_limit(client, project_and_users):
    project_id, _editor, viewer = project_and_users
    rows = [_export_row(), _export_row("BRDP-XL-BIG", rule="x" * (EXCEL_CELL_CHAR_LIMIT + 5))]
    resp = await client.post(f"/api/projects/{project_id}/export.xlsx", json={"rows": rows}, headers=viewer)
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] == "cell_too_large" and detail["limit"] == EXCEL_CELL_CHAR_LIMIT
    assert detail["cells"] == [{"id": "BRDP-XL-BIG", "field": "Rule", "length": EXCEL_CELL_CHAR_LIMIT + 5}]
    assert "BRDP-XL-BIG (Rule)" in detail["message"]


async def test_generic_template_endpoint(client, project_and_users):
    _project_id, _editor, viewer = project_and_users
    assert (await client.get("/api/brdp-template.xlsx")).status_code == 401
    resp = await client.get("/api/brdp-template.xlsx", headers=viewer)
    assert resp.status_code == 200
    assert resp.headers["content-disposition"] == 'attachment; filename="brdp-template.xlsx"'
    assert len(parse_import_file(resp.content, "brdp-template.xlsx")["rows"]) == 10


READ_TEMPLATE_SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "read_template.py"


def test_read_template_script_writes_utf8_even_with_a_cp1252_console(tmp_path):
    """scripts/read_template.py is read by Node (scripts/lib/readXlsx.mjs) as
    UTF-8. On Windows stdout defaults to the console code page (cp1252), and
    every non-ASCII character reached Node as "\ufffd" ("C\ufffddigos" !=
    "Códigos" in verify-xlsx-roundtrip.mjs). PYTHONIOENCODING=cp1252 gives
    the same stdout encoding here."""
    row = {
        "id": "BRDP-UTF8-001",
        "title": "Códigos",
        "definition": "El elemento raíz",
        "proposal": "Campo vacío — sin valor",
        "proposalStatus": "Validated",
        "ruleStatus": "To Do",
        "rule": "",
    }
    path = tmp_path / "utf8.xlsx"
    path.write_bytes(build_export_workbook([row]))
    env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
    out = subprocess.run(
        [sys.executable, str(READ_TEMPLATE_SCRIPT), str(path)], capture_output=True, env=env, check=True
    ).stdout
    rows = json.loads(out.decode("utf-8"))
    assert rows == [
        {
            "ID": "BRDP-UTF8-001",
            "Title": "Códigos",
            "Definition": "El elemento raíz",
            "Proposal": "Campo vacío — sin valor",
            "Proposal Status": "Validated",
            "Rule Status": "To Do",
            "Rule": "",
        }
    ]
